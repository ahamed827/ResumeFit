import io
import re
from pathlib import Path

from app.core.errors import BadRequestError
from app.core.logging import logger


class DocumentParser:
    """
    Extracts text from PDF, DOCX, and TXT files, preserving structure and detecting scanned documents.
    """

    @classmethod
    def extract_text(cls, file_path_or_bytes: Path | bytes, filename: str) -> str:
        ext = Path(filename).suffix.lower()

        if ext == ".pdf":
            return cls._extract_from_pdf(file_path_or_bytes)
        elif ext in [".docx", ".doc"]:
            return cls._extract_from_docx(file_path_or_bytes)
        elif ext == ".txt":
            return cls._extract_from_txt(file_path_or_bytes)
        else:
            raise BadRequestError(f"Unsupported file format: {ext}")

    @classmethod
    def _extract_from_pdf(cls, file_source: Path | bytes) -> str:
        text = ""
        page_count = 0

        # Try PyMuPDF (fitz)
        try:
            import fitz  # PyMuPDF

            if isinstance(file_source, bytes):
                doc = fitz.open(stream=file_source, filetype="pdf")
            else:
                doc = fitz.open(str(file_source))

            page_count = len(doc)
            for page in doc:
                text += page.get_text("text") + "\n\n"
            doc.close()
        except ImportError:
            # Fallback to pypdf
            try:
                import pypdf

                if isinstance(file_source, bytes):
                    reader = pypdf.PdfReader(io.BytesIO(file_source))
                else:
                    reader = pypdf.PdfReader(str(file_source))

                page_count = len(reader.pages)
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n\n"
            except Exception as e:
                logger.error(f"Failed to parse PDF with pypdf: {e}")
                raise BadRequestError(f"Error parsing PDF file: {str(e)}")

        # Detect scanned/empty PDF
        cleaned = text.strip()
        if len(cleaned) < 50:
            raise BadRequestError(
                "Unable to extract selectable text from this PDF. "
                "The file may be a scanned image or flattened document. "
                "Please upload a PDF generated with selectable text or a DOCX document."
            )

        return cls._clean_text(cleaned)

    @classmethod
    def _extract_from_docx(cls, file_source: Path | bytes) -> str:
        try:
            import docx

            if isinstance(file_source, bytes):
                doc = docx.Document(io.BytesIO(file_source))
            else:
                doc = docx.Document(str(file_source))

            paragraphs = []
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text.strip())

            # Also check tables in docx
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        paragraphs.append(" | ".join(row_texts))

            text = "\n\n".join(paragraphs)
            if len(text.strip()) < 30:
                raise BadRequestError("The DOCX file does not contain sufficient readable text.")

            return cls._clean_text(text)
        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
            raise BadRequestError(f"Error reading DOCX file: {str(e)}")

    @classmethod
    def _extract_from_txt(cls, file_source: Path | bytes) -> str:
        try:
            if isinstance(file_source, bytes):
                raw = file_source.decode("utf-8", errors="replace")
            else:
                with open(file_source, "r", encoding="utf-8", errors="replace") as f:
                    raw = f.read()

            if len(raw.strip()) < 20:
                raise BadRequestError("The text file is empty or contains insufficient content.")
            return cls._clean_text(raw)
        except Exception as e:
            raise BadRequestError(f"Error reading text document: {str(e)}")

    @staticmethod
    def _clean_text(text: str) -> str:
        # Normalize non-breaking spaces and redundant whitespaces
        cleaned = text.replace("\xa0", " ")
        cleaned = re.sub(r"[ \t]+", " ", cleaned)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
        return cleaned.strip()


document_parser = DocumentParser()
